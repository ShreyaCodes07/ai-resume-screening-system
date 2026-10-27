from io import BytesIO
from pathlib import Path
import re

import numpy as np
import pymupdf
from docx import Document
from fastapi import FastAPI, UploadFile, File, Form, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.openapi.utils import get_openapi
from sentence_transformers import SentenceTransformer


# --------------------------------------------------
# 1. Application configuration
# --------------------------------------------------

app = FastAPI(title="AI Resume Screening System")

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB per resume
MAX_RESUMES = 20
ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Change these weights to adjust the final score.
SEMANTIC_WEIGHT = 0.60
SKILL_WEIGHT = 0.40

MODEL_NAME = "all-MiniLM-L6-v2"

# Loaded once, not once per uploaded resume.
semantic_model = SentenceTransformer(MODEL_NAME)


app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# --------------------------------------------------
# 2. Swagger file-upload configuration
# --------------------------------------------------

def custom_openapi():
    if app.openapi_schema:
        return app.openapi_schema

    schema = get_openapi(
        title=app.title,
        version="2.0.0",
        description="Semantic and skill-based resume screening",
        routes=app.routes,
    )

    body = schema["components"]["schemas"].get(
        "Body_screen_resumes_screen_resumes_post", {}
    )
    resumes_schema = body.get("properties", {}).get(
        "resumes", {}
    )

    if "items" in resumes_schema:
        resumes_schema["items"]["format"] = "binary"
        resumes_schema["items"].pop("contentMediaType", None)

    app.openapi_schema = schema
    return app.openapi_schema


app.openapi = custom_openapi


# --------------------------------------------------
# 3. Skill dictionary and aliases
# --------------------------------------------------
# Add more skills here as your project grows.
# Each key is the canonical name; values are aliases.

SKILL_ALIASES = {
    "Python": ["python"],
    "Java": ["java"],
    "C++": ["c++", "cpp"],
    "C#": ["c#", "c sharp"],
    "JavaScript": ["javascript", "js"],
    "TypeScript": ["typescript"],
    "SQL": ["sql", "structured query language"],
    "HTML": ["html"],
    "CSS": ["css"],
    "React": ["react", "react.js", "reactjs"],
    "Node.js": ["node.js", "nodejs", "node js"],
    "Express.js": ["express.js", "expressjs", "express"],
    "FastAPI": ["fastapi", "fast api"],
    "Flask": ["flask"],
    "Django": ["django"],
    "REST APIs": ["rest api", "rest apis", "restful api", "restful apis"],
    "GraphQL": ["graphql"],
    "Git": ["git", "git version control"],
    "GitHub": ["github"],
    "Docker": ["docker"],
    "Kubernetes": ["kubernetes", "k8s"],
    "AWS": ["aws", "amazon web services"],
    "Azure": ["azure", "microsoft azure"],
    "GCP": ["gcp", "google cloud platform"],
    "MySQL": ["mysql"],
    "PostgreSQL": ["postgresql", "postgres"],
    "MongoDB": ["mongodb", "mongo db"],
    "SQLite": ["sqlite"],
    "Redis": ["redis"],
    "Pandas": ["pandas"],
    "NumPy": ["numpy"],
    "scikit-learn": ["scikit-learn", "scikit learn", "sklearn"],
    "TensorFlow": ["tensorflow"],
    "PyTorch": ["pytorch"],
    "Machine Learning": ["machine learning", "ml"],
    "Deep Learning": ["deep learning"],
    "NLP": ["nlp", "natural language processing"],
    "Computer Vision": ["computer vision"],
    "Data Analysis": ["data analysis", "data analytics"],
    "Power BI": ["power bi"],
    "Tableau": ["tableau"],
    "Microsoft Excel": ["microsoft excel", "excel"],
    "Statistics": ["statistics", "statistical analysis"],
    "Data Structures": ["data structures", "dsa"],
    "Algorithms": ["algorithms"],
    "Object-Oriented Programming": [
        "object-oriented programming",
        "object oriented programming",
        "oop",
    ],
    "Linux": ["linux"],
    "Agile": ["agile"],
    "Unit Testing": ["unit testing", "unit tests"],
    "CI/CD": ["ci/cd", "continuous integration", "continuous delivery"],
}


# --------------------------------------------------
# 4. Basic endpoints
# --------------------------------------------------

@app.get("/")
def home():
    return {
        "message": "AI Resume Screening System Backend is Running!"
    }


@app.get("/health")
def health_check():
    return {
        "status": "healthy",
        "scoring_method": "Semantic + skill matching",
        "model": MODEL_NAME,
    }


@app.get("/config")
def get_config():
    """Show the active scoring weights and supported skills."""
    return {
        "semantic_weight": SEMANTIC_WEIGHT,
        "skill_weight": SKILL_WEIGHT,
        "model": MODEL_NAME,
        "supported_skills": list(SKILL_ALIASES.keys()),
    }


# --------------------------------------------------
# 5. Safe PDF/DOCX text extraction
# --------------------------------------------------

def extract_resume_text(filename: str, content: bytes) -> str:
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        with pymupdf.open(stream=content, filetype="pdf") as pdf:
            return "\n".join(
                page.get_text() for page in pdf
            ).strip()

    if extension == ".docx":
        document = Document(BytesIO(content))
        parts = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " ".join(cell.text for cell in row.cells)
                )

        return "\n".join(parts).strip()

    raise ValueError("Only PDF and DOCX files are supported.")


# --------------------------------------------------
# 6. Skill extraction
# --------------------------------------------------

def contains_term(text: str, term: str) -> bool:
    """
    Match a term without matching it as part of a longer word.
    Handles punctuation in names such as C++, C#, and CI/CD.
    """
    pattern = (
        r"(?<!\w)"
        + re.escape(term.lower())
        + r"(?!\w)"
    )
    return re.search(pattern, text.lower()) is not None


def extract_skills(text: str) -> list[str]:
    """Return canonical skill names found in text."""
    found = []

    for canonical, aliases in SKILL_ALIASES.items():
        terms = [canonical] + aliases

        if any(contains_term(text, term) for term in terms):
            found.append(canonical)

    return found


# --------------------------------------------------
# 7. Semantic similarity
# --------------------------------------------------

def calculate_semantic_score(
    job_description: str,
    resume_text: str,
) -> float:
    if not job_description.strip() or not resume_text.strip():
        return 0.0

    embeddings = semantic_model.encode(
        [job_description, resume_text],
        normalize_embeddings=True,
        convert_to_numpy=True,
    )

    similarity = float(np.dot(embeddings[0], embeddings[1]))

    # Keep the displayed similarity in the 0-100 range.
    return round(max(0.0, min(1.0, similarity)) * 100, 2)


# --------------------------------------------------
# 8. Explainable combined scoring
# --------------------------------------------------

def analyze_resume(
    job_description: str,
    resume_text: str,
) -> dict:
    # Words and phrases that usually indicate optional skills.
    preferred_markers = [
        "preferred", "preferably", "nice to have", "nice-to-have",
        "bonus", "optional", "plus", "would be a plus",
        "an advantage", "desirable", "good to have", "good-to-have",
        "not mandatory",
    ]

    # Split into lines, semicolon-separated sections, and sentences.
    sections = re.split(r"\n+|;|(?<=[.!?])\s+", job_description)
    required_set = set()
    preferred_set = set()

    for section in sections:
        section_skills = extract_skills(section)
        if not section_skills:
            continue
        section_lower = section.lower()
        is_preferred = any(
            marker in section_lower for marker in preferred_markers
        )
        if is_preferred:
            preferred_set.update(section_skills)
        else:
            required_set.update(section_skills)

    # If a skill appears as both required and preferred, required wins.
    preferred_set -= required_set
    required_skills = [skill for skill in SKILL_ALIASES if skill in required_set]
    preferred_skills = [skill for skill in SKILL_ALIASES if skill in preferred_set]
    resume_skills = extract_skills(resume_text)

    matched_skills = [skill for skill in required_skills if skill in resume_skills]
    missing_skills = [skill for skill in required_skills if skill not in resume_skills]
    preferred_matched_skills = [skill for skill in preferred_skills if skill in resume_skills]
    preferred_missing_skills = [skill for skill in preferred_skills if skill not in resume_skills]

    semantic_score = calculate_semantic_score(job_description, resume_text)

    if required_skills:
        skill_score = round(len(matched_skills) / len(required_skills) * 100, 2)
        final_score = round(
            SEMANTIC_WEIGHT * semantic_score + SKILL_WEIGHT * skill_score,
            2,
        )
        score_explanation = (
            f"{SEMANTIC_WEIGHT:.0%} semantic similarity + "
            f"{SKILL_WEIGHT:.0%} required-skill match. "
            "Preferred skills do not affect the skill score."
        )
    else:
        skill_score = None
        final_score = semantic_score
        if preferred_skills:
            score_explanation = (
                "No required skills were detected. The score uses "
                "semantic similarity alone; preferred skills are reported separately."
            )
        else:
            score_explanation = (
                "No supported skills were detected in the job description, "
                "so the score uses semantic similarity alone. Review or add job requirements."
            )

    # Gather short resume excerpts mentioning matched required or preferred skills.
    evidence = []
    evidence_skills = matched_skills + preferred_matched_skills
    for line in resume_text.splitlines():
        clean_line = line.strip()
        if clean_line and any(
            contains_term(clean_line, skill)
            or any(contains_term(clean_line, alias) for alias in SKILL_ALIASES[skill])
            for skill in evidence_skills
        ):
            evidence.append(clean_line[:300])
        if len(evidence) >= 5:
            break

    return {
        "match_score": final_score,
        "semantic_score": semantic_score,
        "skill_score": skill_score,
        "required_skills": required_skills,
        "preferred_skills": preferred_skills,
        "resume_skills": resume_skills,
        "matched_skills": matched_skills,
        "missing_skills": missing_skills,
        "preferred_matched_skills": preferred_matched_skills,
        "preferred_missing_skills": preferred_missing_skills,
        "skill_match_count": len(matched_skills),
        "required_skill_count": len(required_skills),
        "evidence": evidence,
        "score_explanation": score_explanation,
    }
# --------------------------------------------------
# 9. Main screening endpoint
# --------------------------------------------------

@app.post("/screen-resumes")
async def screen_resumes(
    job_description: str = Form(...),
    resumes: list[UploadFile] = File(...),
):
    job_description = job_description.strip()

    if not job_description:
        raise HTTPException(
            status_code=400,
            detail="Please enter a job description.",
        )

    if len(job_description) > 20000:
        raise HTTPException(
            status_code=400,
            detail="Job description must be 20,000 characters or fewer.",
        )

    if not resumes:
        raise HTTPException(
            status_code=400,
            detail="Please upload at least one resume.",
        )

    if len(resumes) > MAX_RESUMES:
        raise HTTPException(
            status_code=400,
            detail=f"Upload no more than {MAX_RESUMES} resumes at once.",
        )

    results = []

    for resume in resumes:
        filename = resume.filename or "resume"

        try:
            extension = Path(filename).suffix.lower()

            if extension not in ALLOWED_EXTENSIONS:
                raise ValueError("Upload a PDF or DOCX file.")

            # Read at most 10 MB + 1 byte so oversized files
            # can be rejected without reading the entire file.
            content = await resume.read(MAX_FILE_SIZE + 1)

            if len(content) > MAX_FILE_SIZE:
                raise ValueError("File exceeds the 10 MB limit.")

            if not content:
                raise ValueError("The uploaded file is empty.")

            text = extract_resume_text(filename, content)

            if not text.strip():
                raise ValueError(
                    "No readable text found. Scanned PDFs may require OCR."
                )

            if len(text) > 200000:
                raise ValueError(
                    "Extracted resume text is too long to process."
                )

            analysis = analyze_resume(job_description, text)

            # Preserve existing response fields for React.
            results.append({
                "filename": filename,
                "status": "success",
                "match_score": analysis["match_score"],
                "extracted_text": text,
                **analysis,
            })

        except Exception as error:
            # Report a per-file error without losing other results.
            results.append({
                "filename": filename,
                "status": "error",
                "message": str(error)[:500],
            })

        finally:
            await resume.close()

    ranking = sorted(
        [
            item for item in results
            if item["status"] == "success"
        ],
        key=lambda item: item["match_score"],
        reverse=True,
    )

    for rank, item in enumerate(ranking, start=1):
        item["rank"] = rank

    return {
        "job_description": job_description,
        "total_resumes": len(resumes),
        "successful_resumes": len(ranking),
        "results": results,
        "ranking": ranking,
        "scoring_method": "Semantic + skill matching",
        "model": MODEL_NAME,
        "weights": {
            "semantic": SEMANTIC_WEIGHT,
            "skills": SKILL_WEIGHT,
        },
        "notice": (
            "Scores are screening aids, not probabilities of hiring "
            "success. Skill extraction uses a limited dictionary and "
            "may miss synonyms or skills expressed indirectly. Review "
            "the evidence and resumes manually; do not use scores as "
            "the sole basis for hiring decisions."
        ),
    }